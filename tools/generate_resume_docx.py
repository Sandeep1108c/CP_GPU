import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_resume_docx(output_path):
    doc = Document()
    
    # Page setup - Standard A4 with compact 0.28 inch top/bottom and 0.44 inch left/right margins
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.28)
        section.bottom_margin = Inches(0.28)
        section.left_margin = Inches(0.44)
        section.right_margin = Inches(0.44)
    
    # Base font styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(8.9)
    font.color.rgb = RGBColor(30, 30, 30)
    
    # Helper: Set paragraph format
    def set_p_format(p, before=0, after=0.8, line_spacing=1.01, align=WD_ALIGN_PARAGRAPH.LEFT):
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.line_spacing = line_spacing
        p.alignment = align

    # Helper: Add bottom border to section headings
    def add_bottom_border(paragraph, color_hex="1E3A8A", size_pt="6"):
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="{size_pt}" w:space="1" w:color="{color_hex}"/></w:pBdr>')
        pPr.append(pBdr)

    # Helper: Section Heading
    def add_section_heading(title):
        p = doc.add_paragraph()
        set_p_format(p, before=3.5, after=1.8)
        run = p.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(9.8)
        run.font.color.rgb = RGBColor(30, 58, 138) # Deep Navy (#1E3A8A)
        add_bottom_border(p, color_hex="2563EB", size_pt="6")
        return p

    # Helper: Item Header with Left and Right text using tab stop
    def add_item_header(left_title, left_subtitle, right_text, bold_title=True, italic_sub=True):
        p = doc.add_paragraph()
        set_p_format(p, before=1.6, after=0.4)
        
        # 7.39 inches is the right margin position for 8.27 - 0.44 - 0.44
        r_pos = int(7.39 * 1440)
        pPr = p._element.get_or_add_pPr()
        tabs = parse_xml(f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:pos="{r_pos}"/></w:tabs>')
        pPr.append(tabs)
        
        r1 = p.add_run(left_title)
        r1.bold = bold_title
        r1.font.size = Pt(9.1)
        r1.font.color.rgb = RGBColor(15, 23, 42)
        
        if left_subtitle:
            r2 = p.add_run(f" | {left_subtitle}")
            r2.italic = italic_sub
            r2.font.size = Pt(8.4)
            r2.font.color.rgb = RGBColor(71, 85, 105)
            
        r_tab = p.add_run(f"\t{right_text}")
        r_tab.bold = True
        r_tab.font.size = Pt(8.6)
        r_tab.font.color.rgb = RGBColor(30, 41, 59)
        return p

    # Helper: Project overview / summary sentence
    def add_project_summary(summary_text):
        p = doc.add_paragraph()
        set_p_format(p, before=0, after=0.8, line_spacing=1.01)
        r = p.add_run(summary_text)
        r.italic = True
        r.font.size = Pt(8.4)
        r.font.color.rgb = RGBColor(71, 85, 105)
        return p

    # Helper: Bullet point
    def add_bullet(lead_text, body_text):
        p = doc.add_paragraph(style='List Bullet')
        set_p_format(p, before=0, after=0.6, line_spacing=1.01)
        p.paragraph_format.left_indent = Inches(0.18)
        
        if lead_text:
            r_lead = p.add_run(lead_text + " ")
            r_lead.bold = True
            r_lead.font.size = Pt(8.6)
            r_lead.font.color.rgb = RGBColor(15, 23, 42)
            
        r_body = p.add_run(body_text)
        r_body.font.size = Pt(8.6)
        r_body.font.color.rgb = RGBColor(51, 65, 85)
        return p

    # ==========================================
    # HEADER
    # ==========================================
    p_name = doc.add_paragraph()
    set_p_format(p_name, before=0, after=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_name = p_name.add_run("TUTIKA SANDEEP")
    r_name.bold = True
    r_name.font.size = Pt(17.0)
    r_name.font.color.rgb = RGBColor(15, 23, 42)

    p_contact = doc.add_paragraph()
    set_p_format(p_contact, before=0, after=2.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    contact_parts = [
        ("+91-87638-24136", "tel:+918763824136"),
        ("sandeep110809@gmail.com", "mailto:sandeep110809@gmail.com"),
        ("linkedin.com/in/t-sandeep", "https://linkedin.com/in/t-sandeep"),
        ("github.com/Sandeep1108c", "https://github.com/Sandeep1108c"),
        ("Hyderabad, Telangana", None)
    ]
    for i, (text, url) in enumerate(contact_parts):
        r = p_contact.add_run(text)
        r.font.size = Pt(8.3)
        r.font.color.rgb = RGBColor(71, 85, 105)
        if i < len(contact_parts) - 1:
            r_sep = p_contact.add_run("   |   ")
            r_sep.font.size = Pt(8.3)
            r_sep.font.color.rgb = RGBColor(148, 163, 184)

    # ==========================================
    # EDUCATION (Relevant Coursework removed as requested)
    # ==========================================
    add_section_heading("Education")
    
    add_item_header("Woxsen University", "Bachelor of Technology – Computer Science & Engineering (2nd Year)", "2025 – 2029 (Expected)")
    add_item_header("Sankalp Senior Secondary School, Ralaba", "Senior Secondary (XII) – PCM (CBSE)", "2025")
    add_item_header("D.A.V Public School, Berhampur", "Secondary (X) (CBSE)", "2023")

    # ==========================================
    # PATENT & RESEARCH
    # ==========================================
    add_section_heading("Patent & Research")
    
    add_item_header("AquaSense – IoT-Based Smart Water Recycling & Groundwater Monitoring System", "Patent Filed | Prototype Built", "2025")
    add_bullet("Closed-Loop Water Management:", "Designed and implemented an automated IoT water treatment system that purifies greywater, monitors groundwater levels, and controls aquifer recharge using real-time multi-sensor feedback.")
    add_bullet("Dual-Condition Recharge Algorithm:", "Engineered an intelligent recharge algorithm ensuring treated water is only diverted into recharge wells when water quality meets safety standards AND soil moisture indicates capacity, preventing aquifer contamination.")
    add_bullet("Multi-Sensor Instrumentation:", "Integrated 7 sensors (turbidity, pH, TDS, temperature, dual soil moisture, gas detection) with proportional chemical dosing pumps on an ESP32 microcontroller with live telemetry dashboard.")
    add_bullet("Scalable Deployment Target:", "Architected as a cost-effective, decentralized alternative to industrial PLC/SCADA installations for residential communities, farms, and local water conservation.")

    # ==========================================
    # PROJECTS (Ranked: 1. FPGA GPU + nanoGPT, 2. MSME Copilot, 3. Zenith)
    # ==========================================
    add_section_heading("Projects")

    # Project 1: Custom FPGA GPU + Edge nanoGPT AI System
    add_item_header("Custom FPGA-Based 2D GPU & Edge nanoGPT AI System", "Verilog, Gowin EDA, PyTorch, C, ESP32-S3, HDMI", "GitHub")
    add_project_summary("Engineered an end-to-end hardware graphics pipeline on a Gowin FPGA delivering 640×480@60Hz video, integrated with a custom-trained 4.8M nanoGPT edge model for intelligent autonomous agent dialogue.")
    add_bullet("Hardware Rasterization Engines:", "Architected and synthesized a custom 2D GPU processor on a Gowin GW2AR-18C FPGA (20K LUTs), implementing hardware Bresenham line and solid rectangle rasterizers with 18-bit signed coordinate clipping.")
    add_bullet("Hardware Sprite Blitter Core:", "Designed a 2D block transfer engine supporting multi-size sprites (16x16 to 64x64), dual-axis hardware mirroring (flip-X/Y), and zero-overhead color-key transparency.")
    add_bullet("4-Port SDRAM Memory Arbiter:", "Built a prioritized memory arbiter interfacing onboard 64Mbit SDRAM, interleaving 640×480 @ 60 Hz real-time video scanout (Priority 1), SPI asset DMA, blitter operations, and geometry rasterization.")
    add_bullet("Edge nanoGPT LLM Training & SPI Driver:", "Trained a custom 4.8M-parameter nanoGPT transformer in PyTorch for embedded agent dialogue; developed a 20 MHz SPI slave command processor with double-buffering page-flip synchronization.")

    # Project 2: MSME Copilot (Ranked #2)
    add_item_header("MSME Copilot – Business Intelligence & Advisory Platform", "Python, Streamlit, Gemini API, Tavily API, Plotly, FPDF2", "GitHub")
    add_project_summary("Built a multi-agent decision support platform that parses raw financial balance sheets, benchmarks metrics against live industry data, and generates actionable PDF audit roadmaps.")
    add_bullet("Multi-Agent Collaborative Advisory:", "Built a multi-agent decision support system with 5 specialized LLM personas (Orchestrator, Data Analyst, Market Researcher, Financial Planner, Critic) to automate financial diagnosis and operational consulting for small enterprises.")
    add_bullet("Financial Metric Parsing & Visualization:", "Engineered automated data ingestion pipelines parsing unstructured business financial statements into interactive metric summaries, cash-flow trendlines, and comparative Plotly analytics.")
    add_bullet("Market-Verified Advisory & PDF Reports:", "Integrated real-time web search to validate financial assumptions against industry benchmarks, and automated compilation of downloadable executive audit summaries and PDF strategy roadmaps.")

    # Project 3: Zenith (Ranked #3)
    add_item_header("Zenith – Agentic Autonomous Deep Research System", "Python, LangGraph, ChromaDB, Tavily API, Gemini API", "GitHub")
    add_project_summary("Developed an autonomous recursive AI research engine that decomposes ambiguous queries into focused search vectors, executes multi-hop web exploration, and synthesizes verified briefs.")
    add_bullet("Recursive Multi-Hop Investigation:", "Developed an autonomous research agent that recursively decomposes complex research queries into focused search vectors, crawls web sources, identifies knowledge gaps, and loops back to fill missing context.")
    add_bullet("Stateful Agent Orchestration:", "Implemented a LangGraph cyclic state machine enforcing structured exploration trajectories, tool execution safety, dynamic reasoning loops, and deterministic error recovery.")
    add_bullet("Semantic Caching & Synthesis:", "Integrated ChromaDB vector store for semantic caching of search trajectories to eliminate redundant queries, and automated compilation of multi-source findings into structured, cited technical briefs.")

    # ==========================================
    # TECHNICAL SKILLS (Enhanced AI/LLM Tools & Software Frameworks)
    # ==========================================
    add_section_heading("Technical Skills")

    skills_data = [
        ("AI & LLM Frameworks:", "PyTorch, LangGraph, LangChain, Hugging Face Transformers, ChromaDB (Vector DB), RAG Pipelines, Google Gemini API, OpenAI API, Tavily Search API"),
        ("Software & Development Tools:", "Python, Streamlit, Plotly, REST APIs, Git, GitHub, Linux / Bash, Postman, FPDF2, VS Code"),
        ("Hardware & Digital Design:", "Verilog HDL, FPGA Architecture, Gowin EDA, Digital Logic Design, Computer Architecture, Memory Arbiters, SDRAM Controllers, SPI, I2C, UART, HDMI / TMDS"),
        ("Embedded & IoT Systems:", "ESP32 / ESP32-S3, Arduino, PlatformIO, Sensor Interfacing (ADC, I2C, SPI), Actuator & Pump Control, Hardware Debugging"),
        ("Programming Languages:", "Python, C, C++, Verilog, Java, JavaScript, SQL")
    ]

    for title, desc in skills_data:
        p_sk = doc.add_paragraph()
        set_p_format(p_sk, before=0, after=0.5)
        r_t = p_sk.add_run(f"•  {title} ")
        r_t.bold = True
        r_t.font.size = Pt(8.4)
        r_t.font.color.rgb = RGBColor(15, 23, 42)
        
        r_d = p_sk.add_run(desc)
        r_d.font.size = Pt(8.4)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # ==========================================
    # ACHIEVEMENTS & LEADERSHIP (SIH removed, ICRISAT updated to prototype demo)
    # ==========================================
    add_section_heading("Achievements & Leadership")

    add_bullet("Invited Prototype Demonstration at ICRISAT Hyderabad:", "Selected to present and demonstrate the working hardware prototype of patent-pending IoT water recycling & aquifer monitoring system (AquaSense) to agricultural scientists and engineering leadership at ICRISAT headquarters, Hyderabad.")
    add_bullet("Patent Filed:", "Filed Indian Patent for 'IoT-Based Smart Water Recycling and Groundwater Monitoring System' (AquaSense) backed by a fully validated, field-tested physical hardware prototype.")
    add_bullet("Competitive Hackathon Participant:", "Active competitor in premier national and corporate hackathons including AMD Hackathon, IBM Build-a-Thon (BOB), and university engineering hackathons, developing integrated hardware-software prototypes.")

    doc.save(output_path)
    print(f">> Successfully generated resume document at: {output_path}")

if __name__ == '__main__':
    target = r"D:\CP GPU\Tutika_Sandeep_Resume.docx"
    create_resume_docx(target)
